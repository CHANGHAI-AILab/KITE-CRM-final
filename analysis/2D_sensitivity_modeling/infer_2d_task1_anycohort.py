#!/usr/bin/env python3
"""Task 1 2D inference on ANY cohort (Training/Validation/Test cohorts).

Uses the frozen 2D sensitivity DenseNet-121 ensemble (seed0/1/2) trained for the
3D-vs-2D architecture ablation. NOTHING is trained or tuned here: weights are loaded
as-is, the threshold is not touched. Slice rule is the same label-agnostic
central-axial-slice used in training.

Locates each id's kidney clip (Z,H,W .nii.gz) by normalized id across the --img-dirs
(searched in priority order). Directory arguments may be absolute or relative to
sensitivity_paths.DATA_ROOT. The case label table is sensitivity_paths.TASK1_LABELS_XLSX;
checkpoints come from CKDIR_TASK1. Writes per-case probabilities + cohort + whether a
file was found.

Usage:
  python infer_2d_task1_anycohort.py \
     --img-dirs "<abnormal_clip_dir>,<normal_clip_dir>,<test_site_clip_dir>" \
     --cohorts "Test 1,Test 2,Test 3,Test 4,Test 5,Test 6,Training" \
     --out preds_2d_task1_allcohorts.csv
"""
import os, glob, csv, argparse, numpy as np, SimpleITK as sitk, torch, torch.nn as nn
import torchvision.models as M
from scipy.ndimage import zoom

import sensitivity_paths as P

XLSX = P.TASK1_LABELS_XLSX
CKDIR = P.CKDIR_TASK1
DEV = "cuda" if torch.cuda.is_available() else "cpu"
MEAN = np.array([0.485, 0.456, 0.406], np.float32)
STD = np.array([0.229, 0.224, 0.225], np.float32)


def norm(b):
    b = str(b).split('_')[0].split('.')[0].lstrip('0')
    return b or '0'


def read_xlsx_targets(cohorts):
    """id(norm) -> (label, cohort_name, raw_id) for rows whose dataset_name in cohorts."""
    import openpyxl
    ws = openpyxl.load_workbook(XLSX, read_only=True).active
    hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
    bi = hdr.index('编号')
    dni = hdr.index('dataset_name')
    gi = [i for i, h in enumerate(hdr) if h and '任务1' in str(h) and '金标准' in str(h)][0]
    out = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[dni] in cohorts:
            k = norm(r[bi])
            try:
                lab = int(r[gi])
            except (TypeError, ValueError):
                lab = None
            out[k] = (lab, r[dni], r[bi])
    return out


def build_index(img_dirs):
    """norm id -> filepath, first hit wins in priority order."""
    idx = {}
    for d in img_dirs:
        dd = d if os.path.isabs(d) else os.path.join(P.DATA_ROOT, d)
        for f in glob.glob(f"{dd}/*.nii.gz") + glob.glob(f"{dd}/**/*.nii.gz", recursive=True):
            k = norm(os.path.basename(f))
            idx.setdefault(k, f)
    return idx


def central_slice(path):
    a = sitk.GetArrayFromImage(sitk.ReadImage(path)).astype(np.float32) / 255.0
    z = a.shape[0]; c = z // 2
    sl = a[max(0, c - 1):min(z, c + 2)].mean(0)
    sl = zoom(sl, (224 / sl.shape[0], 224 / sl.shape[1]), order=1)
    img = np.stack([sl, sl, sl], 0)
    return ((img - MEAN[:, None, None]) / STD[:, None, None]).astype(np.float32)


def load_models():
    nets = []
    for s in (0, 1, 2):
        m = M.densenet121(weights=None)
        m.classifier = nn.Linear(m.classifier.in_features, 2)
        m.load_state_dict(torch.load(f"{CKDIR}/seed{s}.pt", map_location=DEV))
        nets.append(m.to(DEV).eval())
    return nets


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--img-dirs", required=True)
    ap.add_argument("--cohorts", default="Test 1,Test 2,Test 3,Test 4,Test 5,Test 6,Training")
    ap.add_argument("--gpu", default="0")
    ap.add_argument("--out", default="preds_2d_task1_allcohorts.csv")
    a = ap.parse_args()
    os.environ["CUDA_VISIBLE_DEVICES"] = a.gpu

    cohorts = [c.strip() for c in a.cohorts.split(",")]
    targets = read_xlsx_targets(set(cohorts))
    idx = build_index([d.strip() for d in a.img_dirs.split(",")])
    nets = load_models()
    print(f"targets={len(targets)} indexed_files={len(idx)}", flush=True)

    rows = []; found = 0
    for i, (k, (lab, coh, raw)) in enumerate(sorted(targets.items())):
        path = idx.get(k)
        if path is None:
            rows.append([raw, k, coh, lab, "", "MISSING"]); continue
        try:
            x = torch.from_numpy(central_slice(path))[None].to(DEV)
            ps = []
            with torch.no_grad():
                for m in nets:
                    ps.append(float(torch.softmax(m(x), 1)[0, 1].cpu()))
            rows.append([raw, k, coh, lab, f"{np.mean(ps):.6f}", os.path.basename(path)])
            found += 1
        except Exception as e:
            rows.append([raw, k, coh, lab, "", f"ERR:{type(e).__name__}"])
        if (i + 1) % 100 == 0:
            print(f"  {i+1}/{len(targets)} found={found}", flush=True)

    outp = os.path.join(P.WORK_DIR, a.out)
    with open(outp, "w", newline="") as fh:
        w = csv.writer(fh)
        w.writerow(["id_raw", "id_norm", "cohort", "label", "prob_2d", "file_or_status"])
        w.writerows(rows)
    by = {}
    for r in rows:
        by.setdefault(r[2], [0, 0])
        by[r[2]][0] += 1
        by[r[2]][1] += 1 if r[4] else 0
    print(f"\n[DONE] -> {outp}")
    for c in cohorts:
        if c in by:
            print(f"  {c:<12} total={by[c][0]:>5} inferred={by[c][1]:>5} missing={by[c][0]-by[c][1]:>5}")


if __name__ == "__main__":
    main()
